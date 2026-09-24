"""Read contact details out of an uploaded CV.

Two steps, both deterministic — there is no LLM anywhere in this file:

1. Get the text. PyMuPDF for PDFs, python-docx for .docx, plain decode for
   .txt. A PDF with no text layer (a scan, or a "print to PDF" of an image)
   falls back to OCR when rapidocr-onnxruntime happens to be installed.
2. Pull the identity fields out with regexes. Email, phone, LinkedIn and
   GitHub are all fixed shapes, which is exactly where a regex beats a model:
   no tokens spent, and no chance of a hallucinated handle reaching a matcher
   that links records automatically.

Everything here returns best-effort values for a human to confirm in the form.
Nothing is saved straight from this module.
"""

from __future__ import annotations

import io
import logging
import re

logger = logging.getLogger(__name__)

# Rasterization DPI for the OCR fallback. 200 balances accuracy against
# speed/memory for a CV-sized page.
_OCR_DPI = 200

# Beyond this we are reading someone's thesis, not their CV. The identity
# fields live in the first page or two regardless.
_MAX_TEXT_CHARS = 60_000


# ── Text extraction ──────────────────────────────────────────────────
def extract_text(filename: str, content: bytes) -> tuple[str, str]:
    """Return ``(text, how)`` for an uploaded file.

    ``how`` names the route taken ("pdf", "pdf+ocr", "docx", "text") so the UI
    can tell a recruiter why a scanned CV came back with nothing.

    Raises ValueError for a file type this cannot read.
    """
    lower = (filename or "").lower()

    if lower.endswith(".pdf"):
        text = _extract_pdf(content)
        if _is_thin(text):
            ocr_text = _extract_pdf_ocr(content)
            if len(ocr_text.strip()) > len(text.strip()):
                return ocr_text[:_MAX_TEXT_CHARS], "pdf+ocr"
        return text[:_MAX_TEXT_CHARS], "pdf"

    if lower.endswith(".docx"):
        return _extract_docx(content)[:_MAX_TEXT_CHARS], "docx"

    if lower.endswith((".txt", ".md", ".rtf")):
        return content.decode("utf-8", errors="replace")[:_MAX_TEXT_CHARS], "text"

    if lower.endswith(".doc"):
        raise ValueError(
            "Legacy .doc files cannot be read. Save it as .docx or PDF and upload again."
        )

    raise ValueError("Unsupported file type. Upload a PDF, DOCX or TXT.")


def _is_thin(text: str) -> bool:
    """Too little text to be a real CV — almost certainly a scanned image.

    Deliberately generous: a PDF whose text layer holds only a header still
    reads as "thin", and trying OCR on it costs a few seconds and can only
    improve the result (the longer of the two wins).
    """
    return len(re.sub(r"\s+", "", text or "")) < 120


def _extract_pdf(content: bytes) -> str:
    try:
        import pymupdf  # PyMuPDF >= 1.24 exposes this name
    except ImportError:  # pragma: no cover - older PyMuPDF
        import fitz as pymupdf  # type: ignore[no-redef]

    pages: list[str] = []
    with pymupdf.open(stream=content, filetype="pdf") as doc:
        for page in doc:
            pages.append(page.get_text("text") or "")
    return "\n".join(pages)


def _extract_pdf_ocr(content: bytes) -> str:
    """OCR every page. Returns "" when no OCR engine is installed.

    Optional by design — rapidocr-onnxruntime pulls ~250MB of models, which is
    not worth forcing on an office machine that mostly receives digital CVs.
    """
    try:
        from rapidocr_onnxruntime import RapidOCR
    except ImportError:
        logger.info("Scanned PDF and no OCR engine installed (rapidocr-onnxruntime).")
        return ""

    try:
        import numpy as np

        try:
            import pymupdf
        except ImportError:  # pragma: no cover
            import fitz as pymupdf  # type: ignore[no-redef]

        engine = RapidOCR()
        zoom = _OCR_DPI / 72.0
        out: list[str] = []
        with pymupdf.open(stream=content, filetype="pdf") as doc:
            for page in doc:
                pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom))
                img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(
                    pix.height, pix.width, pix.n
                )
                if pix.n == 4:  # RGBA -> RGB
                    img = img[:, :, :3]
                result, _ = engine(img)
                if result:
                    out.append("\n".join(line[1] for line in result))
        return "\n".join(out)
    except Exception:
        logger.exception("OCR fallback failed.")
        return ""


def _extract_docx(content: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(content))
    parts = [p.text for p in doc.paragraphs]
    # Contact details in a CV are very often laid out in a borderless table,
    # whose text lives outside `paragraphs` entirely.
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


# ── Field extraction ─────────────────────────────────────────────────
EMAIL_RE = re.compile(r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b")
PHONE_RE = re.compile(
    r"(?:(?:\+?\d{1,3}[\s.-]?)?(?:\(?\d{2,4}\)?[\s.-]?)?\d{3,4}[\s.-]?\d{3,4}(?:[\s.-]?\d{1,4})?)"
)
LINKEDIN_RE = re.compile(
    r"(?:https?://)?(?:[A-Za-z]{2,3}\.)?linkedin\.com/(?:in|pub)/[A-Za-z0-9\-_%.]+",
    re.IGNORECASE,
)
GITHUB_RE = re.compile(
    r"(?:https?://)?(?:www\.)?github\.com/[A-Za-z0-9\-_.]+",
    re.IGNORECASE,
)

# A profile slug continued on the next line: lowercase letters, digits and
# hyphens. Requiring lowercase is what stops an ordinary following word
# ("Skills", "Education") from being glued onto the URL.
#
# The leading hyphen is allowed because a wrapped URL frequently breaks AT a
# hyphen in the slug — "…/in/karim" then "-hossain-98765" — and without it
# that slug silently truncates to "karim". The "must contain a letter" rule
# below is what stops a bullet line ("- Led a team") coming through: it
# contributes a bare "-", which carries no letter and is refused.
_SLUG_CONTINUATION_RE = re.compile(r"^-?[a-z0-9][a-z0-9\-_%.]*|^-")

# Characters that can only mean the next line started something NEW — another
# URL, or an email address. A profile slug contains none of them, so whatever
# precedes one is not a continuation however slug-shaped it looks.
_NOT_A_CONTINUATION = frozenset(":/@?#")


def _slug_continuation(token: str, url_so_far: str) -> str:
    """The part of ``token`` continuing a slug split across a line, or "".

    Three rules, all aimed at not swallowing the next line of the CV:

    * must start lowercase (after an optional hyphen) — "Skills",
      "Education", "EXPERIENCE" are out;
    * must contain a letter — "2019" from a "2019 - Present" line is out,
      unless the URL breaks mid-token ("...rahman-"), which only happens when
      the slug really is unfinished;
    * must not be the head of a new URL or an email. CVs stack their contact
      links on consecutive lines, so the line after a LinkedIn URL is very
      often the GitHub one — and "https" or "github.com" or "karim" is
      perfectly slug-shaped right up to the ":" or "/" or "@" that follows it.
    """
    match = _SLUG_CONTINUATION_RE.match(token)
    if not match:
        return ""
    lead = match.group(0)
    rest = token[len(lead) :]
    if rest and rest[0] in _NOT_A_CONTINUATION:
        return ""
    if any(ch.isalpha() for ch in lead):
        return lead
    return lead if url_so_far.endswith(("-", "_", ".")) else ""


def _first_profile_url(pattern: re.Pattern[str], text: str) -> str | None:
    """First profile URL matching ``pattern``, rejoined across a line break.

    PDF extraction splits long URLs at the column edge:

        https://www.linkedin.com/in/abdulla
        h-al-mamun-m-sc-cs

    Without rejoining, the slug is silently TRUNCATED to "abdulla" — and
    because a profile match links two records automatically, two different
    people whose URLs happen to break at the same prefix would be treated as
    one person. Over-joining is the safe direction by comparison: it yields a
    slug that matches nothing, costing a link rather than causing a wrong one.
    """
    match = pattern.search(text or "")
    if not match:
        return None

    url = match.group(0)
    tail = text[match.end() :]
    if tail.startswith(("\n", "\r")):
        continuation = tail.lstrip("\r\n").split(None, 1)
        if continuation:
            url += _slug_continuation(continuation[0], url)

    # PDF extraction commonly glues a trailing separator onto a URL.
    return url.rstrip(".,;:)|").strip() or None


def extract_email(text: str) -> str | None:
    match = EMAIL_RE.search(text or "")
    return match.group(0) if match else None


def extract_phone(text: str) -> str | None:
    if not text:
        return None
    for raw in PHONE_RE.findall(text):
        digits = re.sub(r"\D", "", raw)
        if 7 <= len(digits) <= 15:
            return raw.strip()
    return None


def extract_linkedin(text: str) -> str | None:
    return _first_profile_url(LINKEDIN_RE, text)


def extract_github(text: str) -> str | None:
    return _first_profile_url(GITHUB_RE, text)


# Lines that are section headers, not names.
_NOT_A_NAME = re.compile(
    r"(?i)\b(curriculum\s+vitae|resume|r[ée]sum[ée]|cv|profile|contact|address|"
    r"objective|summary|personal\s+(details|information))\b"
)


def extract_name(text: str) -> str | None:
    """Best guess at the candidate's name: the first line that looks like one.

    A CV puts the name at the top in larger type, which extraction flattens to
    "the first non-empty line". That is right often enough to save typing and
    wrong often enough that the UI must keep it editable — which it does.
    """
    for raw_line in (text or "").splitlines()[:12]:
        line = raw_line.strip(" \t|•·-–—_*#")
        if not (3 <= len(line) <= 60):
            continue
        if "@" in line or _NOT_A_NAME.search(line):
            continue
        if any(ch.isdigit() for ch in line):
            continue
        words = line.split()
        if not (2 <= len(words) <= 5):
            continue
        if all(re.fullmatch(r"[A-Za-z.'\-]+", w) for w in words):
            # PDF extraction pads names out with extra spaces when the
            # original was letter-spaced for effect; a display name reading
            # "Karim  Hossain" looks like a typo nobody made.
            return " ".join(words)
    return None


def parse_cv(filename: str, content: bytes) -> dict:
    """Everything this module can tell you about an uploaded CV."""
    text, how = extract_text(filename, content)
    return {
        "extractedVia": how,
        "textChars": len(text),
        "name": extract_name(text),
        "email": extract_email(text),
        "phone": extract_phone(text),
        "linkedin": extract_linkedin(text),
        "github": extract_github(text),
    }
